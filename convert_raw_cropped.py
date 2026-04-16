import argparse
import csv
from pathlib import Path

import openpyxl
import pydicom
from pydicom.dataset import FileMetaDataset
from pydicom.uid import ExplicitVRLittleEndian, generate_uid


KNOWN_VL_NAMES = ("VL1", "VL2", "VL3")


def parse_args():
    parser = argparse.ArgumentParser(
        description="Crop raw IVUS DICOM cine runs using Excel frame ranges."
    )
    parser.add_argument(
        "--input-dir",
        type=Path,
        required=True,
        help="Directory containing case subfolders with raw DICOM files.",
    )
    parser.add_argument(
        "--output-dir",
        type=Path,
        required=True,
        help="Directory where cropped DICOM files will be written.",
    )
    parser.add_argument(
        "--excel-path",
        type=Path,
        required=True,
        help="Path to Excel file with frame crop parameters.",
    )
    parser.add_argument(
        "--case-start",
        type=int,
        default=None,
        help="Only process cases with numeric IDs greater than or equal to this value.",
    )
    parser.add_argument(
        "--case-end",
        type=int,
        default=None,
        help="Only process cases with numeric IDs less than or equal to this value.",
    )
    parser.add_argument(
        "--study-col",
        type=str,
        default="STUDY_NO.",
        help="Excel column containing the study number.",
    )
    parser.add_argument(
        "--excel-is-1-based",
        action=argparse.BooleanOptionalAction,
        default=True,
        help="Whether Excel frame numbering starts at 1.",
    )
    parser.add_argument(
        "--overwrite",
        action="store_true",
        help="Overwrite existing output files.",
    )
    parser.add_argument(
        "--log-csv",
        type=Path,
        default=None,
        help="Optional CSV path for per-file crop logs.",
    )
    parser.add_argument(
        "--dry-run",
        action="store_true",
        help="Validate inputs and print planned outputs without writing files.",
    )
    args = parser.parse_args()

    if args.case_start is not None and args.case_end is not None and args.case_start > args.case_end:
        parser.error("--case-start cannot be greater than --case-end")

    return args


def case_id_in_range(case_id: str, case_start: int | None, case_end: int | None) -> bool:
    try:
        case_num = int(case_id)
    except ValueError:
        return False

    if case_start is not None and case_num < case_start:
        return False
    if case_end is not None and case_num > case_end:
        return False
    return True


def normalize_case_id(value) -> str | None:
    if value is None:
        return None

    text = str(value).strip()
    if not text:
        return None

    if text.endswith(".0"):
        text = text[:-2]

    try:
        return f"{int(text):03d}"
    except ValueError:
        return None


def load_excel_crop_params(excel_path: Path, study_col: str, excel_is_1_based: bool) -> dict[str, dict[str, tuple[int, int]]]:
    workbook = openpyxl.load_workbook(str(excel_path), data_only=True)
    worksheet = workbook["Sheet1"]

    header = next(worksheet.iter_rows(max_row=1, values_only=True))
    col_idx = {name: i for i, name in enumerate(header)}

    if study_col not in col_idx:
        raise ValueError(f"Column '{study_col}' not found in Excel. Found: {list(col_idx)}")

    vl_columns = []
    for vl_name in KNOWN_VL_NAMES:
        first_col = f"ANN_1STFRAME_{vl_name}"
        last_col = f"ANN_LASTFRAME_{vl_name}"
        if first_col in col_idx and last_col in col_idx:
            vl_columns.append((vl_name, col_idx[first_col], col_idx[last_col]))

    if not vl_columns:
        raise ValueError("No ANN_1STFRAME_/ANN_LASTFRAME_ VL columns found in Excel")

    params: dict[str, dict[str, tuple[int, int]]] = {}
    for row in worksheet.iter_rows(min_row=2, values_only=True):
        case_id = normalize_case_id(row[col_idx[study_col]])
        if case_id is None:
            continue

        case_params = {}
        for vl_name, first_idx, last_idx in vl_columns:
            first_val = row[first_idx]
            last_val = row[last_idx]
            if first_val is None or last_val is None:
                continue

            first_frame = int(first_val)
            last_frame = int(last_val)
            if excel_is_1_based:
                start_idx = first_frame - 1
                end_idx = last_frame - 1
            else:
                start_idx = first_frame
                end_idx = last_frame

            if start_idx > end_idx:
                start_idx, end_idx = end_idx, start_idx

            case_params[vl_name] = (start_idx, end_idx)

        if case_params:
            params[case_id] = case_params

    return params


def is_dicom_file(path: Path) -> bool:
    try:
        pydicom.dcmread(str(path), stop_before_pixels=True, force=True)
        return True
    except Exception:
        return False


def collect_case_folders(input_dir: Path, case_start: int | None, case_end: int | None) -> list[Path]:
    case_dirs = []
    for path in sorted(input_dir.iterdir()):
        if not path.is_dir():
            continue
        if case_id_in_range(path.name, case_start, case_end):
            case_dirs.append(path)
    return case_dirs


def collect_dicom_files(case_dir: Path) -> list[Path]:
    return sorted(path for path in case_dir.rglob("*.dcm") if path.is_file() and is_dicom_file(path))


def infer_vl_name(dicom_path: Path) -> tuple[str, str | None]:
    name_upper = dicom_path.name.upper()
    for vl_name in KNOWN_VL_NAMES:
        if vl_name in name_upper:
            return vl_name, None
    return "VL1", f"[WARNING] {dicom_path}: could not infer VL name from filename, defaulting to VL1"


def get_number_of_frames(dataset) -> int:
    try:
        return int(getattr(dataset, "NumberOfFrames", 1))
    except Exception:
        return 1


def crop_dicom_file(dicom_path: Path, start_idx: int, end_idx: int):
    dataset = pydicom.dcmread(str(dicom_path), force=True)
    total_frames = get_number_of_frames(dataset)

    if start_idx < 0 or end_idx >= total_frames:
        raise ValueError(
            f"Out of bounds. start={start_idx}, end={end_idx}, total_frames={total_frames}"
        )

    pixels = dataset.pixel_array
    if pixels.ndim < 3:
        raise ValueError(f"Unexpected pixel_array shape: {pixels.shape}")

    cropped = pixels[start_idx : end_idx + 1]
    cropped_frames = int(cropped.shape[0])

    output_dataset = dataset.copy()
    output_dataset.PixelData = cropped.tobytes()
    output_dataset.NumberOfFrames = str(cropped_frames)

    if not hasattr(output_dataset, "file_meta") or output_dataset.file_meta is None:
        output_dataset.file_meta = FileMetaDataset()

    output_dataset.file_meta.TransferSyntaxUID = ExplicitVRLittleEndian
    output_dataset.is_little_endian = True
    output_dataset.is_implicit_VR = False
    output_dataset.SOPInstanceUID = generate_uid()
    if hasattr(output_dataset.file_meta, "MediaStorageSOPInstanceUID"):
        output_dataset.file_meta.MediaStorageSOPInstanceUID = output_dataset.SOPInstanceUID

    return output_dataset, total_frames, cropped_frames


def write_log_csv(log_path: Path, rows: list[dict]):
    fieldnames = [
        "case",
        "vl_name",
        "status",
        "reason",
        "input_path",
        "output_path",
        "total_frames_original",
        "start_frame_excel",
        "end_frame_excel",
        "start_index_python",
        "end_index_python",
        "saved_frames",
    ]
    log_path.parent.mkdir(parents=True, exist_ok=True)
    with log_path.open("w", newline="", encoding="utf-8") as file_obj:
        writer = csv.DictWriter(file_obj, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)


def main():
    args = parse_args()

    if not args.input_dir.exists():
        raise FileNotFoundError(f"Input directory not found: {args.input_dir}")
    if not args.excel_path.exists():
        raise FileNotFoundError(f"Excel file not found: {args.excel_path}")

    crop_params = load_excel_crop_params(args.excel_path, args.study_col, args.excel_is_1_based)
    case_dirs = collect_case_folders(args.input_dir, args.case_start, args.case_end)
    if not case_dirs:
        raise RuntimeError("No case folders found for the requested case range")

    if not args.dry_run:
        args.output_dir.mkdir(parents=True, exist_ok=True)

    print(f"Loaded crop parameters for {len(crop_params)} cases from {args.excel_path}")
    print(f"Found {len(case_dirs)} case folder(s) under {args.input_dir}")

    log_rows = []
    success = 0
    failed = 0

    for case_dir in case_dirs:
        case_id = case_dir.name
        dicom_files = collect_dicom_files(case_dir)

        if not dicom_files:
            failed += 1
            message = f"[ERROR] No DICOM files found in {case_dir}"
            print(message)
            log_rows.append(
                {
                    "case": case_id,
                    "vl_name": "",
                    "status": "FAILED",
                    "reason": "No DICOM files found",
                    "input_path": "",
                    "output_path": "",
                    "total_frames_original": "",
                    "start_frame_excel": "",
                    "end_frame_excel": "",
                    "start_index_python": "",
                    "end_index_python": "",
                    "saved_frames": "",
                }
            )
            continue

        if len(dicom_files) > 1:
            print(f"[WARNING] Study {case_id}: found {len(dicom_files)} DICOM files. Using the first one.")

        dicom_path = dicom_files[0]
        vl_name, warning_message = infer_vl_name(dicom_path)
        if warning_message:
            print(warning_message)

        if case_id not in crop_params:
            failed += 1
            print(f"[ERROR] No crop params found for case {case_id}")
            log_rows.append(
                {
                    "case": case_id,
                    "vl_name": vl_name,
                    "status": "FAILED",
                    "reason": "No crop params found for case",
                    "input_path": str(dicom_path),
                    "output_path": "",
                    "total_frames_original": "",
                    "start_frame_excel": "",
                    "end_frame_excel": "",
                    "start_index_python": "",
                    "end_index_python": "",
                    "saved_frames": "",
                }
            )
            continue

        case_crop_params = crop_params[case_id]
        if vl_name not in case_crop_params:
            failed += 1
            print(f"[ERROR] No crop params found for {case_id}/{vl_name}")
            log_rows.append(
                {
                    "case": case_id,
                    "vl_name": vl_name,
                    "status": "FAILED",
                    "reason": "No crop params found for pullback",
                    "input_path": str(dicom_path),
                    "output_path": "",
                    "total_frames_original": "",
                    "start_frame_excel": "",
                    "end_frame_excel": "",
                    "start_index_python": "",
                    "end_index_python": "",
                    "saved_frames": "",
                }
            )
            continue

        start_idx, end_idx = case_crop_params[vl_name]
        expected_frames = end_idx - start_idx + 1
        out_path = args.output_dir / case_id / f"{case_id}_{vl_name}_cropped.dcm"

        if out_path.exists() and not args.overwrite and not args.dry_run:
            failed += 1
            print(f"[ERROR] Output already exists and --overwrite was not set: {out_path}")
            log_rows.append(
                {
                    "case": case_id,
                    "vl_name": vl_name,
                    "status": "FAILED",
                    "reason": "Output exists",
                    "input_path": str(dicom_path),
                    "output_path": str(out_path),
                    "total_frames_original": "",
                    "start_frame_excel": start_idx + 1 if args.excel_is_1_based else start_idx,
                    "end_frame_excel": end_idx + 1 if args.excel_is_1_based else end_idx,
                    "start_index_python": start_idx,
                    "end_index_python": end_idx,
                    "saved_frames": "",
                }
            )
            continue

        try:
            output_dataset, total_frames, cropped_frames = crop_dicom_file(dicom_path, start_idx, end_idx)
            if cropped_frames != expected_frames:
                raise RuntimeError(
                    f"Saved frame count {cropped_frames} did not match expected count {expected_frames}"
                )

            if args.dry_run:
                print(
                    f"[DRY RUN] {case_id}/{vl_name}: total_frames={total_frames}, "
                    f"saved_frames={cropped_frames} -> {out_path}"
                )
            else:
                out_path.parent.mkdir(parents=True, exist_ok=True)
                output_dataset.save_as(str(out_path), write_like_original=False)
                print(
                    f"[OK] {case_id}/{vl_name}: total_frames={total_frames}, "
                    f"saved_frames={cropped_frames} -> {out_path}"
                )

            success += 1
            log_rows.append(
                {
                    "case": case_id,
                    "vl_name": vl_name,
                    "status": "SUCCESS",
                    "reason": "",
                    "input_path": str(dicom_path),
                    "output_path": str(out_path),
                    "total_frames_original": total_frames,
                    "start_frame_excel": start_idx + 1 if args.excel_is_1_based else start_idx,
                    "end_frame_excel": end_idx + 1 if args.excel_is_1_based else end_idx,
                    "start_index_python": start_idx,
                    "end_index_python": end_idx,
                    "saved_frames": cropped_frames,
                }
            )
        except Exception as exc:
            failed += 1
            print(f"[ERROR] {case_id}/{vl_name}: {exc}")
            log_rows.append(
                {
                    "case": case_id,
                    "vl_name": vl_name,
                    "status": "FAILED",
                    "reason": str(exc),
                    "input_path": str(dicom_path),
                    "output_path": str(out_path),
                    "total_frames_original": "",
                    "start_frame_excel": start_idx + 1 if args.excel_is_1_based else start_idx,
                    "end_frame_excel": end_idx + 1 if args.excel_is_1_based else end_idx,
                    "start_index_python": start_idx,
                    "end_index_python": end_idx,
                    "saved_frames": "",
                }
            )

    print("\nCropping complete")
    print(f"Successful: {success}")
    print(f"Failed: {failed}")

    if args.log_csv:
        write_log_csv(args.log_csv, log_rows)
        print(f"Log saved to: {args.log_csv}")


if __name__ == "__main__":
    main()
